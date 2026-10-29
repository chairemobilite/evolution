# Copyright 2026, Polytechnique Montreal and contributors
# This file is licensed under the MIT License.
# License text available at https://opensource.org/licenses/MIT

# Note: Tests for scripts/generate_survey.py's Excel integrity/loading entry points.

import os
import shutil
from unittest.mock import MagicMock

import openpyxl
import pytest  # pyright: ignore[reportMissingImports]
from openpyxl.cell.read_only import EMPTY_CELL
from openpyxl.styles import PatternFill

from helpers.generator_helpers import (
    delete_file_if_exists,
)
from tests.mocked_excel import add_mocked_excel_sheet, create_mocked_excel_data
import scripts.generate_survey as generate_survey_module
from scripts.generate_survey import (
    check_excel_integrity,
    check_excel_integrity_with_messages,
    check_excel_integrity_with_warnings,
    generate_survey,
)
from survey_definition.excel_survey_definition_reader import (
    EMPTY_ROWS_WARNING_THRESHOLD,
    ExcelSurveyDefinitionReader,
)

MOCKED_EXCEL_FILE = "src/tests/references/test.xlsx"
DEMO_EXCEL_FILE = os.path.join(
    os.path.dirname(__file__),
    "../../../../example/demo_generator/references/Household_Travel_Generate_Survey.xlsx",
)

VALID_SECTIONS_HEADERS = [
    "section",
    "title_fr",
    "title_en",
    "in_nav",
    "template",
    "parent_section",
    "abbreviation",
]
VALID_SECTIONS_ROWS = [["home", "Accueil", "Home", True, "", "", "HM_"]]

VALID_CONDITIONALS_HEADERS = [
    "conditional_name",
    "logical_operator",
    "path",
    "comparison_operator",
    "value",
    "parentheses",
    "value_when_hidden",
]
VALID_CONDITIONALS_ROWS = [["cond1", "", "household.size", "===", 1, "", ""]]


VALID_WIDGETS_HEADERS = ["questionName", "label::fr", "label::en"]


def _create_workbook(
    sections_rows=VALID_SECTIONS_ROWS,
    conditionals_rows=VALID_CONDITIONALS_ROWS,
    widgets_headers=VALID_WIDGETS_HEADERS,
    sections_headers=VALID_SECTIONS_HEADERS,
    choices_rows=(),
):
    """Write a workbook whose Widgets sheet holds the survey languages (its label::<language> columns)."""
    workbook = create_mocked_excel_data("Sections", sections_headers, sections_rows)
    add_mocked_excel_sheet(
        workbook, "Conditionals", VALID_CONDITIONALS_HEADERS, conditionals_rows
    )
    add_mocked_excel_sheet(workbook, "Widgets", widgets_headers, [])
    add_mocked_excel_sheet(workbook, "Labels", ["namespace", "key"], [])
    add_mocked_excel_sheet(
        workbook,
        "Choices",
        [
            "choicesName",
            "value",
            "label::fr",
            "label::en",
            "spreadChoicesName",
            "conditional",
        ],
        list(choices_rows),
    )


class TestCheckExcelIntegrityWithMessages:
    def test_valid_file_returns_true(self):
        _create_workbook()
        try:
            assert check_excel_integrity_with_messages(MOCKED_EXCEL_FILE) == (True, [])
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_invalid_file_returns_false(self):
        ok, messages = check_excel_integrity_with_messages("nonexistent.xlsx")
        assert ok is False
        assert len(messages) == 1

    def test_invalid_rows_in_both_sheets_return_every_message(self):
        _create_workbook(
            sections_rows=[["home", "Accueil", "Home", True, "", "", ""]],
            conditionals_rows=[["cond1", "", "", "===", 1, "", ""]],
        )
        try:
            ok, messages = check_excel_integrity_with_messages(MOCKED_EXCEL_FILE)
            assert ok is False
            # Conditionals first: the tables are validated in dependency order
            assert messages == [
                "Error in Conditionals - Required field is missing in row 2. Missing fields: ['path']",
                "Error in Sections - Required field is missing in row 2. Missing fields: ['abbreviation']",
            ]
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_missing_title_is_reported_with_the_other_sheets_errors(self):
        """The rule that needs the survey languages runs with each sheet's other checks."""
        _create_workbook(
            sections_rows=[["home", "", "Home", True, "", "", "HM_"]],
            conditionals_rows=[["cond1", "", "", "===", 1, "", ""]],
        )
        try:
            assert check_excel_integrity_with_messages(MOCKED_EXCEL_FILE) == (
                False,
                [
                    "Error in Conditionals - Required field is missing in row 2. Missing fields: ['path']",
                    "Error in Sections - Invalid row 2: title_fr is required when in_nav is true",
                ],
            )
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_checks_the_languages_of_the_widgets_sheet(self):
        # The Widgets sheet has label::es, but the Sections sheet has no title_es column
        _create_workbook(widgets_headers=["questionName", "label::fr", "label::es"])
        try:
            assert check_excel_integrity_with_messages(MOCKED_EXCEL_FILE) == (
                False,
                [
                    "Error in Sections - Invalid row 2: title_es is required when in_nav is true"
                ],
            )
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_reports_a_title_column_in_another_language(self):
        """A title_<language> column for a language of no label:: column, e.g. a typo, is reported."""
        _create_workbook(
            widgets_headers=["questionName", "label::fr"],
        )
        try:
            ok, messages = check_excel_integrity_with_messages(MOCKED_EXCEL_FILE)
            assert ok is False
            assert messages == [
                "Error in Sections - Unexpected columns: ['title_en']. The survey "
                "languages are ['fr']."
            ]
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_reports_a_widgets_sheet_without_label_columns(self):
        _create_workbook(widgets_headers=["questionName", "label"])
        try:
            ok, messages = check_excel_integrity_with_messages(MOCKED_EXCEL_FILE)
            assert ok is False
            assert messages == [
                "Missing label::<language> columns in Widgets sheet: the survey "
                "languages are read from them, e.g. label::fr and label::en."
            ]
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_reports_invalid_languages(self):
        _create_workbook(widgets_headers=["questionName", "label::French"])
        try:
            ok, messages = check_excel_integrity_with_messages(MOCKED_EXCEL_FILE)
            assert ok is False
            assert messages == [
                "Invalid survey languages ['French']: each must be a language code "
                "like 'fr' or 'en-CA'."
            ]
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)


class TestCheckExcelIntegrityChoices:
    def test_reports_a_choice_conditional_missing_from_the_conditionals_sheet(self):
        _create_workbook(
            choices_rows=[["yesNo", "yes", "Oui", "Yes", "", "nopeConditional"]]
        )
        try:
            assert check_excel_integrity_with_messages(MOCKED_EXCEL_FILE) == (
                False,
                [
                    "Error in Choices - Invalid row 2: conditional 'nopeConditional' "
                    "does not match any conditional_name of the Conditionals table"
                ],
            )
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_reports_the_choices_problems_with_the_other_sheets_ones(self):
        _create_workbook(
            sections_rows=[["home", "Accueil", "Home", True, "", "", ""]],
            choices_rows=[["yesNo", "yes", "Oui", "", "", ""]],
        )
        try:
            ok, messages = check_excel_integrity_with_messages(MOCKED_EXCEL_FILE)
            assert ok is False
            assert messages == [
                "Error in Sections - Required field is missing in row 2. Missing fields: ['abbreviation']",
                "Error in Choices - Invalid row 2: label::en is required",
            ]
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)


class TestCheckExcelIntegrity:
    def test_valid_file_returns_true(self):
        _create_workbook()
        try:
            assert check_excel_integrity(MOCKED_EXCEL_FILE) is True
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_invalid_file_returns_false(self):
        assert check_excel_integrity("nonexistent.xlsx") is False


class TestCheckExcelIntegrityWithWarnings:
    def test_warns_about_many_empty_rows_after_the_data(self):
        _create_workbook()
        # A formatted cell far below the data, with no value
        workbook = openpyxl.load_workbook(MOCKED_EXCEL_FILE)
        workbook["Widgets"].cell(
            row=EMPTY_ROWS_WARNING_THRESHOLD + 1, column=1
        ).fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        workbook.save(MOCKED_EXCEL_FILE)
        try:
            assert check_excel_integrity_with_warnings(MOCKED_EXCEL_FILE) == (
                True,
                [],
                [
                    f"Widgets sheet has {EMPTY_ROWS_WARNING_THRESHOLD} empty "
                    "rows after its last row with data (row 1). Delete them in Excel "
                    "to make the file faster to open."
                ],
            )
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_unreadable_file_has_no_warnings(self):
        ok, messages, warnings = check_excel_integrity_with_warnings("nonexistent.xlsx")
        assert ok is False
        assert len(messages) == 1
        assert warnings == []


class TestGenerateSurvey:
    def test_invalid_excel_file_aborts_generation(self, tmp_path):
        config_path = tmp_path / "generatorConfigs.yaml"
        config_path.write_text(
            f"survey_folder_path: {tmp_path}\n"
            "excel_file_path: nonexistent.xlsx\n"
            "enabled_scripts: {}\n"
        )
        with pytest.raises(Exception, match="Excel integrity check failed"):
            generate_survey(str(config_path))
        # Nothing was generated in the survey folder
        assert list(tmp_path.iterdir()) == [config_path]

    def test_creates_a_locales_folder_per_survey_language(self, tmp_path):
        _create_workbook(
            widgets_headers=["questionName", "label::en"],
            sections_headers=[h for h in VALID_SECTIONS_HEADERS if h != "title_fr"],
            sections_rows=[["home", "Home", True, "", "", "HM_"]],
        )
        config_path = tmp_path / "generatorConfigs.yaml"
        config_path.write_text(
            f"survey_folder_path: {tmp_path}\n"
            f"excel_file_path: {MOCKED_EXCEL_FILE}\n"
            "enabled_scripts: {}\n"
        )
        try:
            generate_survey(str(config_path))
            assert sorted(os.listdir(tmp_path / "locales")) == ["en"]
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_loads_the_survey_definition_once(self, tmp_path, monkeypatch):
        load_survey_definition = MagicMock(
            wraps=generate_survey_module.load_survey_definition
        )
        monkeypatch.setattr(
            generate_survey_module, "load_survey_definition", load_survey_definition
        )
        _create_workbook()
        config_path = tmp_path / "generatorConfigs.yaml"
        config_path.write_text(
            f"survey_folder_path: {tmp_path}\n"
            f"excel_file_path: {MOCKED_EXCEL_FILE}\n"
        )
        try:
            generate_survey(str(config_path))
            load_survey_definition.assert_called_once()
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_generates_the_choices_from_the_survey_definition(self, tmp_path):
        _create_workbook(choices_rows=[["yesNo", "yes", "Oui", "Yes", "", ""]])
        config_path = tmp_path / "generatorConfigs.yaml"
        config_path.write_text(
            f"survey_folder_path: {tmp_path}\n"
            f"excel_file_path: {MOCKED_EXCEL_FILE}\n"
        )
        try:
            generate_survey(str(config_path), only_scripts={"generate_choices"})
            choices_ts = (
                tmp_path / "src" / "survey" / "common" / "choices.tsx"
            ).read_text(encoding="utf-8")
            assert "export const yesNo: ChoiceType[] = [" in choices_ts
            assert "yes: Oui" in (
                tmp_path / "locales" / "fr" / "choices.yaml"
            ).read_text(encoding="utf-8")
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_reads_the_excel_file_once_for_every_script(self, tmp_path, monkeypatch):
        if not os.path.isfile(DEMO_EXCEL_FILE):
            pytest.skip(f"Demo generator file not found: {DEMO_EXCEL_FILE}")
        # The CSV copy is written next to the Excel file: copy it out of the repo.
        excel_file_path = tmp_path / "references" / "survey.xlsx"
        excel_file_path.parent.mkdir()
        shutil.copy(DEMO_EXCEL_FILE, excel_file_path)
        load_workbook = MagicMock(wraps=openpyxl.load_workbook)
        monkeypatch.setattr(openpyxl, "load_workbook", load_workbook)
        enabled_scripts = "".join(
            f"    {script_key}: {script_key != 'generate_excel'}\n"
            for script_key in generate_survey_module.SUPPORTED_SCRIPT_KEYS
        )
        config_path = tmp_path / "generatorConfigs.yaml"
        config_path.write_text(
            f"survey_folder_path: {tmp_path}\n"
            f"excel_file_path: {excel_file_path}\n"
            f"enabled_scripts:\n{enabled_scripts}"
        )
        generate_survey(str(config_path))
        load_workbook.assert_called_once()

    def test_empty_rows_between_data_rows_generate_the_same_files(
        self, tmp_path, monkeypatch
    ):
        """An empty row left between rows of the sheets not yet in the survey definition is skipped."""
        if not os.path.isfile(DEMO_EXCEL_FILE):
            pytest.skip(f"Demo generator file not found: {DEMO_EXCEL_FILE}")
        enabled_scripts = "".join(
            f"    {script_key}: {script_key not in ('generate_excel', 'copy_excel_to_csv')}\n"
            for script_key in generate_survey_module.SUPPORTED_SCRIPT_KEYS
        )

        def generate(survey_folder_path):
            survey_folder_path.mkdir()
            config_path = tmp_path / f"{survey_folder_path.name}.yaml"
            config_path.write_text(
                f"survey_folder_path: {survey_folder_path}\n"
                f"excel_file_path: {DEMO_EXCEL_FILE}\n"
                f"enabled_scripts:\n{enabled_scripts}"
            )
            generate_survey(str(config_path))
            return {
                path.relative_to(survey_folder_path): path.read_bytes()
                for path in sorted(survey_folder_path.rglob("*"))
                if path.is_file()
            }

        expected_files = generate(tmp_path / "without_empty_rows")

        class ReaderWithEmptyRows(ExcelSurveyDefinitionReader):
            def __init__(self, excel_file_path: str) -> None:
                super().__init__(excel_file_path)
                for sheet_name in ["Widgets", "InputRange"]:
                    # The sheets are private to the reader: add the empty row there.
                    rows = self._sheets[sheet_name].rows
                    rows.insert(2, (EMPTY_CELL,) * len(rows[0]))

        monkeypatch.setattr(
            generate_survey_module, "ExcelSurveyDefinitionReader", ReaderWithEmptyRows
        )
        assert generate(tmp_path / "with_empty_rows") == expected_files
