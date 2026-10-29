# Copyright 2026, Polytechnique Montreal and contributors
# This file is licensed under the MIT License.
# License text available at https://opensource.org/licenses/MIT

# Note: Tests for survey_definition/excel_survey_definition_reader.py, which reads the
# Excel file once and returns its tables as dicts.

import openpyxl
import pytest  # pyright: ignore[reportMissingImports]
from openpyxl.styles import PatternFill

from survey_definition.excel_survey_definition_reader import (
    EMPTY_COLUMNS_WARNING_THRESHOLD,
    EMPTY_ROWS_WARNING_THRESHOLD,
    ExcelSurveyDefinitionReader,
)

HEADERS = ["section", "in_nav"]
DATA_ROWS = [["home", True], ["profile", False]]


def _save_workbook(path, styled_cells=(), rows=DATA_ROWS) -> str:
    """
    Save a one-sheet workbook ("Sections") with HEADERS and `rows`, plus a fill on
    each (row, column) of `styled_cells`: a formatted cell with no value, as Excel
    keeps when a cell outside the data is formatted.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sections"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    for row_number, column_number in styled_cells:
        sheet.cell(row=row_number, column=column_number).fill = PatternFill(
            fill_type="solid", fgColor="FFFF00"
        )
    workbook.save(path)
    return str(path)


def _read_sections(path: str) -> list[dict]:
    return ExcelSurveyDefinitionReader(path).read_table("Sections", HEADERS)


EXPECTED_TABLE = [
    {"section": "home", "in_nav": True},
    {"section": "profile", "in_nav": False},
]


@pytest.mark.parametrize(
    "styled_cells",
    [
        pytest.param([], id="no styled cell"),
        pytest.param([(10, 1)], id="styled cell below the data"),
        pytest.param([(10, 1), (11, 1)], id="styled cells below the data"),
        pytest.param([(100_000, 1)], id="styled cell far below the data"),
        pytest.param([(1, 50)], id="styled cell right of the data"),
    ],
)
def test_reads_only_the_area_with_data(tmp_path, styled_cells):
    path = _save_workbook(tmp_path / "survey.xlsx", styled_cells=styled_cells)
    assert _read_sections(path) == EXPECTED_TABLE


def test_keeps_empty_rows_between_rows_with_data(tmp_path):
    """An empty row in the middle is kept, so the table's checks still report it."""
    path = _save_workbook(
        tmp_path / "survey.xlsx", rows=[["home", True], [None, None], ["end", None]]
    )
    assert _read_sections(path) == [
        {"section": "home", "in_nav": True},
        {"section": None, "in_nav": None},
        {"section": "end", "in_nav": None},
    ]


def test_reads_a_sheet_with_only_a_styled_empty_row_as_an_empty_table(tmp_path):
    path = _save_workbook(tmp_path / "survey.xlsx", styled_cells=[(2, 1)], rows=[])
    assert _read_sections(path) == []


def test_pads_short_rows_to_the_width_of_the_sheet(tmp_path):
    path = _save_workbook(tmp_path / "survey.xlsx", rows=[["home"], ["end", True]])
    assert _read_sections(path) == [
        {"section": "home", "in_nav": None},
        {"section": "end", "in_nav": True},
    ]


def test_missing_table_raises(tmp_path):
    reader = ExcelSurveyDefinitionReader(_save_workbook(tmp_path / "survey.xlsx"))

    with pytest.raises(Exception, match="Sheet with name Widgets does not exist"):
        reader.read_table("Widgets", [])


def test_missing_required_field_raises(tmp_path):
    reader = ExcelSurveyDefinitionReader(_save_workbook(tmp_path / "survey.xlsx"))

    with pytest.raises(
        Exception, match="Missing expected header in Sections sheet: abbreviation"
    ):
        reader.read_table("Sections", [*HEADERS, "abbreviation"])


def test_rejects_a_file_that_is_not_an_excel_file():
    with pytest.raises(Exception, match="Invalid input file extension for survey.csv"):
        ExcelSurveyDefinitionReader("survey.csv")


def test_lists_every_sheet_in_order(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active.title = "Widgets"
    workbook.create_sheet("Choices")
    workbook.save(tmp_path / "survey.xlsx")
    assert ExcelSurveyDefinitionReader(str(tmp_path / "survey.xlsx")).sheet_names == [
        "Widgets",
        "Choices",
    ]


@pytest.mark.parametrize(
    "rows, expected_values",
    [
        pytest.param(DATA_ROWS, DATA_ROWS, id="no empty row"),
        pytest.param(
            [["home", True], [None, None], ["end", False]],
            [["home", True], ["end", False]],
            id="empty row between rows with data",
        ),
        pytest.param(
            [[None, True]], [[None, True]], id="row with a value in any column is kept"
        ),
        pytest.param([], [], id="header only"),
    ],
)
def test_get_data_rows_skips_the_empty_rows(tmp_path, rows, expected_values):
    reader = ExcelSurveyDefinitionReader(
        _save_workbook(tmp_path / "survey.xlsx", rows=rows)
    )
    assert [
        [cell.value for cell in row] for row in reader.get_data_rows("Sections")
    ] == expected_values


ROWS_WARNING = (
    f"Sections sheet has {EMPTY_ROWS_WARNING_THRESHOLD} empty rows after "
    "its last row with data (row 3). Delete them in Excel to make the file faster "
    "to open."
)
COLUMNS_WARNING = (
    f"Sections sheet has {EMPTY_COLUMNS_WARNING_THRESHOLD} empty columns "
    "after its last column with data (column B). Delete them in Excel to make the "
    "file faster to open."
)


# The data fills rows 1 to 3 and columns A and B: a styled cell at row 3 + n or at
# column 2 + n leaves n empty rows or columns after the data.
@pytest.mark.parametrize(
    "styled_cells, expected_warnings",
    [
        pytest.param([(10, 1), (1, 10)], [], id="few empty rows and columns"),
        pytest.param(
            [(3 + EMPTY_ROWS_WARNING_THRESHOLD - 1, 1)],
            [],
            id="rows just under the threshold",
        ),
        pytest.param(
            [(3 + EMPTY_ROWS_WARNING_THRESHOLD, 1)],
            [ROWS_WARNING],
            id="rows at the threshold",
        ),
        pytest.param(
            [(1, 2 + EMPTY_COLUMNS_WARNING_THRESHOLD - 1)],
            [],
            id="columns just under the threshold",
        ),
        pytest.param(
            [(1, 2 + EMPTY_COLUMNS_WARNING_THRESHOLD)],
            [COLUMNS_WARNING],
            id="columns at the threshold",
        ),
        pytest.param(
            # A styled cell in an empty row after the data still widens the sheet.
            [(3 + EMPTY_ROWS_WARNING_THRESHOLD, 2 + EMPTY_COLUMNS_WARNING_THRESHOLD)],
            [ROWS_WARNING, COLUMNS_WARNING],
            id="rows and columns at the threshold",
        ),
    ],
)
def test_warns_about_many_empty_rows_or_columns_after_the_data(
    tmp_path, styled_cells, expected_warnings
):
    path = _save_workbook(tmp_path / "survey.xlsx", styled_cells=styled_cells)
    assert ExcelSurveyDefinitionReader(path).warnings == expected_warnings
