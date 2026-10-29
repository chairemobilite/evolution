# Copyright 2026, Polytechnique Montreal and contributors
# This file is licensed under the MIT License.
# License text available at https://opensource.org/licenses/MIT

# Note: Tests for helpers/excel_workbook.py, which reads the Excel file once.

import openpyxl
import pytest  # pyright: ignore[reportMissingImports]
from openpyxl.styles import PatternFill

from helpers.excel_workbook import (
    EMPTY_COLUMNS_WARNING_THRESHOLD,
    EMPTY_ROWS_WARNING_THRESHOLD,
    load_excel_workbook,
)

HEADERS = ["section", "in_nav"]
DATA_ROWS = [["home", True], ["profile", False]]


def _save_workbook(path, styled_cells=(), rows=DATA_ROWS) -> str:
    """
    Save a one-sheet workbook ("Sections") with HEADERS and `rows`, plus a fill on
    each (row, column) of `styled_cells`: a formatted cell with no value, as Excel
    keeps when a cell outside the data is formatted.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sections"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    for row_number, column_number in styled_cells:
        sheet.cell(row=row_number, column=column_number).fill = PatternFill(
            fill_type="solid", fgColor="FFFF00"
        )
    workbook.save(path)
    return str(path)


def _values(sheet) -> list[list]:
    return [[cell.value for cell in row] for row in sheet.rows]


EXPECTED_VALUES = [HEADERS, *DATA_ROWS]


@pytest.mark.parametrize(
    "styled_cells, expected_values",
    [
        pytest.param([], EXPECTED_VALUES, id="no styled cell"),
        pytest.param([(10, 1)], EXPECTED_VALUES, id="styled cell below the data"),
        pytest.param(
            [(100_000, 1)], EXPECTED_VALUES, id="styled cell far below the data"
        ),
        pytest.param([(1, 50)], EXPECTED_VALUES, id="styled cell right of the data"),
    ],
)
def test_reads_only_the_area_with_data(tmp_path, styled_cells, expected_values):
    path = _save_workbook(tmp_path / "survey.xlsx", styled_cells=styled_cells)
    assert _values(load_excel_workbook(path)["Sections"]) == expected_values


def test_keeps_empty_rows_between_rows_with_data(tmp_path):
    path = _save_workbook(
        tmp_path / "survey.xlsx", rows=[["home", True], [None, None], ["end", None]]
    )
    assert _values(load_excel_workbook(path)["Sections"]) == [
        HEADERS,
        ["home", True],
        [None, None],
        ["end", None],
    ]


def test_pads_short_rows_to_the_width_of_the_sheet(tmp_path):
    path = _save_workbook(tmp_path / "survey.xlsx", rows=[["home"], ["end", True]])
    assert _values(load_excel_workbook(path)["Sections"]) == [
        HEADERS,
        ["home", None],
        ["end", True],
    ]


def test_reads_every_sheet_in_order(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active.title = "Widgets"
    workbook.create_sheet("Choices")
    workbook.save(tmp_path / "survey.xlsx")
    assert load_excel_workbook(str(tmp_path / "survey.xlsx")).sheetnames == [
        "Widgets",
        "Choices",
    ]


def test_missing_sheet_raises_a_key_error_naming_the_sheet(tmp_path):
    workbook = load_excel_workbook(_save_workbook(tmp_path / "survey.xlsx"))

    with pytest.raises(KeyError, match="Worksheet Widgets does not exist."):
        workbook["Widgets"]


ROWS_WARNING = (
    f"Sections sheet has {EMPTY_ROWS_WARNING_THRESHOLD} empty rows after "
    "its last row with data (row 3). Delete them in Excel to make the file faster "
    "to open."
)
COLUMNS_WARNING = (
    f"Sections sheet has {EMPTY_COLUMNS_WARNING_THRESHOLD} empty columns "
    "after its last column with data (column B). Delete them in Excel to make the "
    "file faster to open."
)


# The data fills rows 1 to 3 and columns A and B: a styled cell at row 3 + n or at
# column 2 + n leaves n empty rows or columns after the data.
@pytest.mark.parametrize(
    "styled_cells, expected_warnings",
    [
        pytest.param([(10, 1), (1, 10)], [], id="few empty rows and columns"),
        pytest.param(
            [(3 + EMPTY_ROWS_WARNING_THRESHOLD - 1, 1)],
            [],
            id="rows just under the threshold",
        ),
        pytest.param(
            [(3 + EMPTY_ROWS_WARNING_THRESHOLD, 1)],
            [ROWS_WARNING],
            id="rows at the threshold",
        ),
        pytest.param(
            [(1, 2 + EMPTY_COLUMNS_WARNING_THRESHOLD - 1)],
            [],
            id="columns just under the threshold",
        ),
        pytest.param(
            [(1, 2 + EMPTY_COLUMNS_WARNING_THRESHOLD)],
            [COLUMNS_WARNING],
            id="columns at the threshold",
        ),
        pytest.param(
            # A styled cell in an empty row after the data still widens the sheet.
            [(3 + EMPTY_ROWS_WARNING_THRESHOLD, 2 + EMPTY_COLUMNS_WARNING_THRESHOLD)],
            [ROWS_WARNING, COLUMNS_WARNING],
            id="rows and columns at the threshold",
        ),
    ],
)
def test_warns_about_many_empty_rows_or_columns_after_the_data(
    tmp_path, styled_cells, expected_warnings
):
    path = _save_workbook(tmp_path / "survey.xlsx", styled_cells=styled_cells)
    assert load_excel_workbook(path).warnings == expected_warnings
