# Copyright 2026, Polytechnique Montreal and contributors
# This file is licensed under the MIT License.
# License text available at https://opensource.org/licenses/MIT

import pytest  # pyright: ignore[reportMissingImports]
from openpyxl.styles import PatternFill
from pydantic import ValidationError

from helpers.generator_helpers import (
    INDENT,
    _read_sheet_rows,
    delete_file_if_exists,
    generate_label_typescript_with_context,
    get_data_rows,
    get_label_context_flags,
    get_workbook,
    load_survey_definition,
)
from tests.mocked_excel import add_mocked_excel_sheet, create_mocked_excel_data

MOCKED_EXCEL_FILE = "src/tests/references/test.xlsx"


class TestGetLabelContextFlags:
    def test_returns_all_false_when_no_tokens_present(self):
        assert get_label_context_flags(labels=["Bonjour", "Hello"]) == (
            False,
            False,
            False,
            False,
        )

    def test_detects_nickname_token(self):
        assert get_label_context_flags(labels=["Bonjour {{nickname}}"]) == (
            True,
            False,
            False,
            False,
        )

    def test_detects_count_token_even_with_spaces(self):
        assert get_label_context_flags(labels=["Hello {{ count }}"]) == (
            False,
            True,
            False,
            False,
        )

    def test_detects_gender_token_with_or_without_spaces_around_colon(self):
        assert get_label_context_flags(labels=["Étudian{{gender:t/te/t·e}}"]) == (
            False,
            False,
            True,
            False,
        )
        assert get_label_context_flags(labels=["Étudian{{gender : t/te/t·e}}"]) == (
            False,
            False,
            True,
            False,
        )

    def test_detects_label_one_presence_and_ignores_whitespace_only(self):
        assert get_label_context_flags(labels_one=["Salut"]) == (
            False,
            False,
            False,
            True,
        )
        assert get_label_context_flags(labels_one=["   "]) == (
            False,
            False,
            False,
            False,
        )

    def test_ignores_missing_labels(self):
        """A label missing in a language is None."""
        assert get_label_context_flags(
            labels=[None, "Hi {{nickname}}"], labels_one=[None, None]
        ) == (
            True,
            False,
            False,
            False,
        )

    def test_detects_all_flags_together(self):
        assert get_label_context_flags(
            labels=["Bonjour **{{nickname}}** ({{count}}) Étudian{{gender:t/te/t·e}}"],
            labels_one=["Hi"],
        ) == (
            True,
            True,
            True,
            True,
        )


class TestGenerateLabelTypescriptWithContext:
    def test_generates_simple_t_function_when_no_context_needed(self):
        result = generate_label_typescript_with_context(
            property_name="label",
            translation_key="sectionA:foo.bar",
            base_indent=INDENT,
            has_nickname=False,
            has_count=False,
            has_gender_context=False,
            has_label_one=False,
        )
        assert result == f"{INDENT}label: (t: TFunction) => t('sectionA:foo.bar')"

    def test_generates_full_function_with_nickname_context(self):
        result = generate_label_typescript_with_context(
            property_name="label",
            translation_key="sectionA:foo.bar",
            base_indent=INDENT,
            has_nickname=True,
            has_count=False,
            has_gender_context=False,
            has_label_one=False,
        )
        assert f"{INDENT}label: (t: TFunction, interview, path) => {{" in result
        assert (
            f"{INDENT}{INDENT}const activePerson = odSurveyHelpers.getPerson({{ interview, path }});"
            in result
        )
        assert (
            f"{INDENT}{INDENT}const nickname = _escape(activePerson?.nickname || t('survey:noNickname'));"
            in result
        )
        assert f"{INDENT}{INDENT}return t('sectionA:foo.bar', {{" in result
        assert f"{INDENT}{INDENT}{INDENT}nickname," in result

    def test_generates_count_persons_when_label_one_is_present(self):
        result = generate_label_typescript_with_context(
            property_name="label",
            translation_key="sectionB:baz",
            base_indent=INDENT,
            has_nickname=False,
            has_count=False,
            has_gender_context=False,
            has_label_one=True,
        )
        assert (
            f"{INDENT}{INDENT}const countPersons = odSurveyHelpers.countPersons({{ interview }});"
            in result
        )
        assert f"{INDENT}{INDENT}{INDENT}count: countPersons," in result

    def test_uses_provided_gender_context_expression(self):
        result = generate_label_typescript_with_context(
            property_name="text",
            translation_key="sectionC:qux",
            base_indent="",
            has_nickname=False,
            has_count=False,
            has_gender_context=True,
            has_label_one=False,
            gender_context_expression="activePerson?.gender",
        )
        assert "text: (t: TFunction, interview, path) => {" in result
        assert (
            "const activePerson = odSurveyHelpers.getPerson({ interview, path });"
            in result
        )
        assert "context: activePerson?.gender," in result


class _Cell:
    def __init__(self, value):
        self.value = value


def _row(*values):
    return tuple(_Cell(value) for value in values)


@pytest.mark.parametrize(
    "rows, expected_values",
    [
        pytest.param(
            [_row("name"), _row("a"), _row("b")], [["a"], ["b"]], id="no empty row"
        ),
        pytest.param(
            [_row("name", "label"), _row("a", "A"), _row(None, None), _row("b", "B")],
            [["a", "A"], ["b", "B"]],
            id="empty row between rows with data",
        ),
        pytest.param(
            [_row("name", "label"), _row(None, "A")],
            [[None, "A"]],
            id="row with a value in any column is kept",
        ),
        pytest.param([_row("name")], [], id="header only"),
    ],
)
def test_get_data_rows(rows, expected_values):
    assert [
        [cell.value for cell in row] for row in get_data_rows(rows)
    ] == expected_values


class TestReadSheetRows:
    """
    Tests for _read_sheet_rows, in particular the trailing blank rows the workbook
    drops when it is read (see ExcelSheet).

    A row with no cell value and no style isn't written to the saved .xlsx at all,
    so it can't reproduce the bug the trimming guards against. A row
    that has a style but no value (e.g. from selecting extra rows and applying
    formatting in a spreadsheet editor) IS written, and openpyxl reads it back as a
    row of all-None values — that's what these tests build, via a styled blank row.
    """

    HEADERS = ["section", "in_nav", "abbreviation"]

    def _style_blank_row(self, workbook, row_number: int) -> None:
        """Give a row a fill (no cell value) and re-save, so it survives the reload."""
        sheet = workbook["Sections"]
        for column in range(1, len(self.HEADERS) + 1):
            sheet.cell(row=row_number, column=column).fill = PatternFill(
                fill_type="solid", fgColor="FFFF00"
            )
        workbook.save(MOCKED_EXCEL_FILE)

    def _read(self) -> list[dict]:
        try:
            return _read_sheet_rows(
                get_workbook(MOCKED_EXCEL_FILE),
                "Sections",
                self.HEADERS,
            )
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_keeps_rows_when_none_are_blank(self):
        create_mocked_excel_data(
            "Sections", self.HEADERS, [["home", True, "HM_"], ["profile", True, "PR_"]]
        )
        rows = self._read()
        assert [row["section"] for row in rows] == ["home", "profile"]

    def test_drops_a_trailing_styled_blank_row(self):
        workbook = create_mocked_excel_data(
            "Sections", self.HEADERS, [["home", True, "HM_"]]
        )
        self._style_blank_row(workbook, row_number=3)
        rows = self._read()
        assert [row["section"] for row in rows] == ["home"]

    def test_drops_more_than_one_trailing_styled_blank_row(self):
        workbook = create_mocked_excel_data(
            "Sections", self.HEADERS, [["home", True, "HM_"]]
        )
        self._style_blank_row(workbook, row_number=3)
        self._style_blank_row(workbook, row_number=4)
        rows = self._read()
        assert [row["section"] for row in rows] == ["home"]

    def test_keeps_an_interior_styled_blank_row(self):
        """Only trailing blank rows are dropped; a blank row in the middle is kept
        (and still invalid, so Sections validation still reports it)."""
        workbook = create_mocked_excel_data(
            "Sections", self.HEADERS, [["home", True, "HM_"]]
        )
        self._style_blank_row(workbook, row_number=3)
        workbook["Sections"].append(["profile", True, "PR_"])
        workbook.save(MOCKED_EXCEL_FILE)
        rows = self._read()
        assert [row["section"] for row in rows] == ["home", None, "profile"]

    def test_drops_a_lone_styled_blank_row(self):
        """A sheet with only a styled blank row reads like a sheet with only headers."""
        workbook = create_mocked_excel_data("Sections", self.HEADERS, [])
        self._style_blank_row(workbook, row_number=2)
        rows = self._read()
        assert rows == []


class TestLoadSurveyDefinition:
    """Tests for load_survey_definition (builds a SurveyDefinition from the Excel file's sheets)."""

    VALID_SECTIONS_HEADERS = [
        "section",
        "title_fr",
        "title_en",
        "in_nav",
        "template",
        "parent_section",
        "abbreviation",
    ]
    VALID_SECTIONS_ROWS = [["home", "Accueil", "Home", True, "", "", "HM_"]]

    VALID_CONDITIONALS_HEADERS = [
        "conditional_name",
        "logical_operator",
        "path",
        "comparison_operator",
        "value",
        "parentheses",
        "value_when_hidden",
    ]
    VALID_CONDITIONALS_ROWS = [["cond1", "", "household.size", "===", 1, "", ""]]

    VALID_LABELS_HEADERS = ["namespace", "key", "label::fr", "label::en"]
    VALID_LABELS_ROWS = [["home", "welcome", "Bienvenue", "Welcome"]]

    VALID_CHOICES_HEADERS = [
        "choicesName",
        "value",
        "label::fr",
        "label::en",
        "spreadChoicesName",
        "conditional",
    ]
    VALID_CHOICES_ROWS = [["yesNo", "yes", "Oui", "Yes", "", ""]]

    def _create_workbook(self, sheet_name, headers, rows):
        """Like create_mocked_excel_data, plus a Widgets sheet holding the survey languages (fr, en), a Labels sheet and a Choices sheet."""
        workbook = create_mocked_excel_data(sheet_name, headers, rows)
        add_mocked_excel_sheet(
            workbook, "Widgets", ["questionName", "label::fr", "label::en"], []
        )
        add_mocked_excel_sheet(
            workbook, "Labels", self.VALID_LABELS_HEADERS, self.VALID_LABELS_ROWS
        )
        add_mocked_excel_sheet(
            workbook, "Choices", self.VALID_CHOICES_HEADERS, self.VALID_CHOICES_ROWS
        )
        return workbook

    def test_valid_file_returns_a_survey_definition(self):
        workbook = self._create_workbook(
            "Sections", self.VALID_SECTIONS_HEADERS, self.VALID_SECTIONS_ROWS
        )
        add_mocked_excel_sheet(
            workbook,
            "Conditionals",
            self.VALID_CONDITIONALS_HEADERS,
            self.VALID_CONDITIONALS_ROWS,
        )
        try:
            survey_definition = load_survey_definition(get_workbook(MOCKED_EXCEL_FILE))
            assert [section.section for section in survey_definition.sections] == [
                "home"
            ]
            assert [
                conditional.conditional_name
                for conditional in survey_definition.conditionals
            ] == ["cond1"]
            assert [
                (label.namespace, label.key, label.label)
                for label in survey_definition.labels
            ] == [("home", "welcome", {"fr": "Bienvenue", "en": "Welcome"})]
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_sections_sheet_ignores_trailing_blank_row(self):
        """A trailing row with no value but a style (e.g. from formatting applied past
        the real data in a spreadsheet editor) is dropped, not reported as a row
        missing every required field. A plain unstyled blank row wouldn't reproduce
        this: openpyxl doesn't even persist it, so the style is what makes the row
        survive the save/reload round trip as a row of None values."""
        workbook = self._create_workbook(
            "Sections", self.VALID_SECTIONS_HEADERS, self.VALID_SECTIONS_ROWS
        )
        sections_sheet = workbook["Sections"]
        for column in range(1, len(self.VALID_SECTIONS_HEADERS) + 1):
            sections_sheet.cell(row=3, column=column).fill = PatternFill(
                fill_type="solid", fgColor="FFFF00"
            )
        add_mocked_excel_sheet(
            workbook,
            "Conditionals",
            self.VALID_CONDITIONALS_HEADERS,
            self.VALID_CONDITIONALS_ROWS,
        )
        try:
            survey_definition = load_survey_definition(get_workbook(MOCKED_EXCEL_FILE))
            assert [section.section for section in survey_definition.sections] == [
                "home"
            ]
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_conditionals_sheet_ignores_trailing_blank_row(self):
        """Same as the Sections sheet: a styled trailing blank row is dropped."""
        workbook = self._create_workbook(
            "Sections", self.VALID_SECTIONS_HEADERS, self.VALID_SECTIONS_ROWS
        )
        add_mocked_excel_sheet(
            workbook,
            "Conditionals",
            self.VALID_CONDITIONALS_HEADERS,
            self.VALID_CONDITIONALS_ROWS,
        )
        conditionals_sheet = workbook["Conditionals"]
        for column in range(1, len(self.VALID_CONDITIONALS_HEADERS) + 1):
            conditionals_sheet.cell(row=3, column=column).fill = PatternFill(
                fill_type="solid", fgColor="FFFF00"
            )
        workbook.save(MOCKED_EXCEL_FILE)
        try:
            survey_definition = load_survey_definition(get_workbook(MOCKED_EXCEL_FILE))
            assert [
                conditional.conditional_name
                for conditional in survey_definition.conditionals
            ] == ["cond1"]
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_conditionals_sheet_without_value_when_hidden_column_is_valid(self):
        """value_when_hidden is optional: its column doesn't need to exist at all."""
        workbook = self._create_workbook(
            "Sections", self.VALID_SECTIONS_HEADERS, self.VALID_SECTIONS_ROWS
        )
        add_mocked_excel_sheet(
            workbook,
            "Conditionals",
            self.VALID_CONDITIONALS_HEADERS[:-1],  # drop value_when_hidden
            [row[:-1] for row in self.VALID_CONDITIONALS_ROWS],
        )
        try:
            survey_definition = load_survey_definition(get_workbook(MOCKED_EXCEL_FILE))
            assert survey_definition.conditionals[0].value_when_hidden is None
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_missing_sections_sheet_raises(self):
        self._create_workbook(
            "OtherSheet",
            self.VALID_SECTIONS_HEADERS,
            self.VALID_SECTIONS_ROWS,
        )
        try:
            with pytest.raises(
                Exception, match="Sheet with name Sections does not exist"
            ):
                load_survey_definition(get_workbook(MOCKED_EXCEL_FILE))
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_missing_conditionals_sheet_raises(self):
        self._create_workbook(
            "Sections",
            self.VALID_SECTIONS_HEADERS,
            self.VALID_SECTIONS_ROWS,
        )
        try:
            with pytest.raises(
                Exception, match="Sheet with name Conditionals does not exist"
            ):
                load_survey_definition(get_workbook(MOCKED_EXCEL_FILE))
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_missing_labels_sheet_raises(self):
        workbook = self._create_workbook(
            "Sections", self.VALID_SECTIONS_HEADERS, self.VALID_SECTIONS_ROWS
        )
        add_mocked_excel_sheet(
            workbook,
            "Conditionals",
            self.VALID_CONDITIONALS_HEADERS,
            self.VALID_CONDITIONALS_ROWS,
        )
        del workbook["Labels"]
        workbook.save(MOCKED_EXCEL_FILE)
        try:
            with pytest.raises(
                Exception, match="Sheet with name Labels does not exist"
            ):
                load_survey_definition(get_workbook(MOCKED_EXCEL_FILE))
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_labels_sheet_missing_a_required_header_raises(self):
        workbook = self._create_workbook(
            "Sections", self.VALID_SECTIONS_HEADERS, self.VALID_SECTIONS_ROWS
        )
        add_mocked_excel_sheet(
            workbook,
            "Conditionals",
            self.VALID_CONDITIONALS_HEADERS,
            self.VALID_CONDITIONALS_ROWS,
        )
        del workbook["Labels"]
        add_mocked_excel_sheet(
            workbook, "Labels", ["namespace", "label::fr", "label::en"], []
        )
        try:
            with pytest.raises(Exception, match="Labels"):
                load_survey_definition(get_workbook(MOCKED_EXCEL_FILE))
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_invalid_sections_row_raises_validation_error(self):
        workbook = self._create_workbook(
            "Sections",
            self.VALID_SECTIONS_HEADERS,
            [[None, "Accueil", "Home", True, "", "", "HM_"]],
        )
        add_mocked_excel_sheet(
            workbook,
            "Conditionals",
            self.VALID_CONDITIONALS_HEADERS,
            self.VALID_CONDITIONALS_ROWS,
        )
        try:
            with pytest.raises(ValidationError):
                load_survey_definition(get_workbook(MOCKED_EXCEL_FILE))
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_invalid_conditionals_row_raises_validation_error(self):
        workbook = self._create_workbook(
            "Sections", self.VALID_SECTIONS_HEADERS, self.VALID_SECTIONS_ROWS
        )
        add_mocked_excel_sheet(
            workbook,
            "Conditionals",
            self.VALID_CONDITIONALS_HEADERS,
            [[None, "", "household.size", "===", 1, "", ""]],
        )
        try:
            with pytest.raises(ValidationError):
                load_survey_definition(get_workbook(MOCKED_EXCEL_FILE))
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_reads_the_survey_languages_from_the_widgets_sheet(self):
        workbook = create_mocked_excel_data(
            "Sections",
            [
                "section",
                "title_es",
                "title_en",
                "in_nav",
                "template",
                "parent_section",
                "abbreviation",
            ],
            [["home", "Inicio", "Home", True, "", "", "HM_"]],
        )
        add_mocked_excel_sheet(
            workbook,
            "Conditionals",
            self.VALID_CONDITIONALS_HEADERS,
            self.VALID_CONDITIONALS_ROWS,
        )
        add_mocked_excel_sheet(
            workbook, "Widgets", ["questionName", "label::es", "label::en"], []
        )
        add_mocked_excel_sheet(workbook, "Labels", self.VALID_LABELS_HEADERS, [])
        add_mocked_excel_sheet(workbook, "Choices", self.VALID_CHOICES_HEADERS, [])
        try:
            survey_definition = load_survey_definition(get_workbook(MOCKED_EXCEL_FILE))
            assert survey_definition.languages == ("es", "en")
            assert survey_definition.sections[0].title == {"es": "Inicio", "en": "Home"}
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_missing_widgets_sheet_raises(self):
        workbook = create_mocked_excel_data(
            "Sections", self.VALID_SECTIONS_HEADERS, self.VALID_SECTIONS_ROWS
        )
        add_mocked_excel_sheet(
            workbook,
            "Conditionals",
            self.VALID_CONDITIONALS_HEADERS,
            self.VALID_CONDITIONALS_ROWS,
        )
        try:
            with pytest.raises(
                Exception, match="Sheet with name Widgets does not exist"
            ):
                load_survey_definition(get_workbook(MOCKED_EXCEL_FILE))
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def _create_valid_workbook(self):
        workbook = self._create_workbook(
            "Sections", self.VALID_SECTIONS_HEADERS, self.VALID_SECTIONS_ROWS
        )
        add_mocked_excel_sheet(
            workbook,
            "Conditionals",
            self.VALID_CONDITIONALS_HEADERS,
            self.VALID_CONDITIONALS_ROWS,
        )
        return workbook

    def test_reads_the_choices_sheet(self):
        workbook = self._create_valid_workbook()
        del workbook["Choices"]
        add_mocked_excel_sheet(
            workbook,
            "Choices",
            self.VALID_CHOICES_HEADERS + ["label_one::fr", "label_one::en", "hidden"],
            [
                ["yes", "yes", "Oui", "Yes", "", "", "", "", True],
                ["yesNo", "", "", "", "yes", "", "", "", ""],
            ],
        )
        try:
            yes, spread = load_survey_definition(
                get_workbook(MOCKED_EXCEL_FILE)
            ).choices
            assert yes.label == {"fr": "Oui", "en": "Yes"}
            assert yes.hidden is True
            assert spread.spread_choices_name == "yes"
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_missing_choices_sheet_raises(self):
        workbook = self._create_valid_workbook()
        del workbook["Choices"]
        workbook.save(MOCKED_EXCEL_FILE)
        try:
            with pytest.raises(
                Exception, match="Sheet with name Choices does not exist"
            ):
                load_survey_definition(get_workbook(MOCKED_EXCEL_FILE))
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)

    def test_choices_sheet_missing_a_required_header_raises(self):
        workbook = self._create_valid_workbook()
        del workbook["Choices"]
        add_mocked_excel_sheet(
            workbook,
            "Choices",
            [h for h in self.VALID_CHOICES_HEADERS if h != "spreadChoicesName"],
            [],
        )
        try:
            with pytest.raises(
                Exception,
                match="Missing expected header in Choices sheet: spreadChoicesName",
            ):
                load_survey_definition(get_workbook(MOCKED_EXCEL_FILE))
        finally:
            delete_file_if_exists(MOCKED_EXCEL_FILE)
