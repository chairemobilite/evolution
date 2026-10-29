# Copyright 2026, Polytechnique Montreal and contributors
# This file is licensed under the MIT License.
# License text available at https://opensource.org/licenses/MIT

# Note: This module reads the survey definition from the Generator Excel file. The file
# is read once, so every script works on the same in-memory copy instead of
# re-opening it. Sheets, rows and cells are details of this reader: the survey
# definition is read as tables of dicts (see ExcelSurveyDefinitionReader).
from typing import List, Tuple

import openpyxl
from openpyxl.cell.read_only import EMPTY_CELL
from openpyxl.utils import get_column_letter

# Number of empty rows after the last row with data, and of empty columns after the
# last column with data, from which a sheet gets a warning. Excel keeps such rows and
# columns when a cell outside the data is formatted, and they make the file slow to
# open for every tool.
EMPTY_ROWS_WARNING_THRESHOLD = 1000
EMPTY_COLUMNS_WARNING_THRESHOLD = 100

# Prefix of the translated columns the survey languages are read from, in the Widgets
# sheet: one column per language, e.g. label::fr, label::en.
LANGUAGE_COLUMN_PREFIX = "label::"


class _ExcelSheet:
    """
    One sheet read into memory, cut to the area that has data: rows after the last row
    with a value, and columns after the last column with a value, are dropped. Every
    row has the same number of cells, each with a `.value` (None when empty).
    """

    def __init__(self, title: str, rows: List[Tuple]) -> None:
        self.title = title
        self.rows = rows


class ExcelSurveyDefinitionReader:
    """
    Reads the survey definition from the Generator Excel file: the survey languages
    (read_languages) and each table, one sheet per table (read_table).

    The whole file is read when the reader is created. `warnings` lists the problems
    found while reading that don't stop the generation, e.g. a sheet with many empty
    rows or columns after its data.
    """

    def __init__(self, excel_file_path: str) -> None:
        _check_excel_file_extension(excel_file_path)
        sheets, self.warnings = _read_workbook(excel_file_path)
        self._sheets = {sheet.title: sheet for sheet in sheets}

    def read_languages(self) -> tuple[str, ...]:
        """
        Read the survey languages from the Widgets sheet: one per `label::<language>`
        column, in column order, e.g. ("fr", "en") for label::fr and label::en.

        The file carries its own languages, so a survey is checked the same way from
        any app or config. Raises if the sheet has no such column.
        """
        headers = [cell.value for cell in self._get_sheet("Widgets").rows[0]]
        languages = tuple(
            header.removeprefix(LANGUAGE_COLUMN_PREFIX)
            for header in headers
            if isinstance(header, str) and header.startswith(LANGUAGE_COLUMN_PREFIX)
        )
        if not languages:
            raise Exception(
                "Missing label::<language> columns in Widgets sheet: the survey languages "
                "are read from them, e.g. label::fr and label::en."
            )
        return languages

    def read_table(self, table_name: str, required_fields: List[str]) -> List[dict]:
        """
        Read every row of the `table_name` sheet (row 1 is the header) as a dict keyed
        by its headers. Raises if the sheet or one of `required_fields` is missing.

        Empty rows between rows with data are kept (still invalid if the table requires
        values); there are no empty rows after the data (see _ExcelSheet).
        """
        headers = self._get_headers(table_name, required_fields)
        return [
            dict(zip(headers, (cell.value for cell in row), strict=True))
            for row in self._get_sheet(table_name).rows[1:]
        ]

    # TODO: Remove the methods below once every sheet is migrated to the survey
    # definition (#2035): they give the scripts that still read the other sheets
    # (e.g. Widgets, InputRange) and the CSV copy access to the sheets' rows.

    @property
    def sheet_names(self) -> List[str]:
        """The names of the file's sheets, in file order."""
        return list(self._sheets)

    def get_rows_and_headers(self, sheet_name: str) -> Tuple[List, List[str]]:
        """
        Return the rows of the sheet, header row included, and its headers: the
        header row's values up to its first empty cell.

        Raises if the sheet is missing or a header has spaces.
        """
        try:
            rows: List = self._get_sheet(sheet_name).rows

            # Filter out None values from headers
            headers = []
            for cell in rows[0]:
                if cell.value is not None:
                    headers.append(cell.value)
                else:
                    # If we find an empty header, stop reading columns here
                    break

            # Error when header has spaces
            if any(" " in str(header) for header in headers):
                raise Exception("Header has spaces")

            return rows, headers

        except Exception as e:
            print(f"Error reading Excel in {sheet_name} sheet: {e}")
            raise e

    def get_headers(self, sheet_name: str, expected_headers: List[str]) -> List:
        """
        Return the values of the sheet's header row. Raises if the sheet or one of
        `expected_headers` is missing.
        """
        return self._get_headers(sheet_name, expected_headers)

    def get_data_rows(self, sheet_name: str) -> List:
        """
        Return the rows of the sheet after the header row, without the empty ones (e.g.
        a blank row left to separate groups of rows), so scripts only see rows with
        data.
        """
        return [row for row in self._get_sheet(sheet_name).rows[1:] if _has_value(row)]

    def _get_headers(self, sheet_name: str, expected_headers: List[str]) -> List:
        current_headers = [cell.value for cell in self._get_sheet(sheet_name).rows[0]]

        # Check if the headers are valid
        for expected in expected_headers:
            if not expected in current_headers:
                raise Exception(
                    f"Missing expected header in {sheet_name} sheet: {expected}"
                )

        return current_headers

    def _get_sheet(self, sheet_name: str) -> _ExcelSheet:
        if sheet_name not in self._sheets:
            raise Exception(f"Sheet with name {sheet_name} does not exist")
        return self._sheets[sheet_name]


def _check_excel_file_extension(file: str) -> None:
    file_lower = file.lower()
    if not (
        file_lower.endswith(".xlsx")
        or file_lower.endswith(".xlsm")
        or file_lower.endswith(".xltx")
        or file_lower.endswith(".xltm")
    ):
        raise Exception(
            f"Invalid input file extension for {file} : must be an Excel 2010+ .xlsx, .xlsm, .xltx or .xltm file"
        )


def _has_value(row: Tuple) -> bool:
    return any(cell.value is not None for cell in row)


def _read_sheet(worksheet) -> Tuple[_ExcelSheet, int, int]:
    """
    Read a read-only worksheet into an _ExcelSheet.

    Returns the sheet and the number of empty rows and empty columns found after its
    last row and last column with data.
    """
    # The dimension saved in the file covers every formatted cell, so reading up to it
    # can mean a million empty rows: read only the rows the file actually has instead.
    worksheet.reset_dimensions()
    rows: List[Tuple] = []
    empty_rows_after_data = 0
    # Each row read ends at its last cell in the file, with or without a value.
    file_width = 0
    for row in worksheet.iter_rows():
        file_width = max(file_width, len(row))
        if _has_value(row):
            # Keep the empty rows between rows with data: they are part of the sheet.
            rows.extend([()] * empty_rows_after_data)
            empty_rows_after_data = 0
            rows.append(row)
        else:
            empty_rows_after_data += 1

    width = max(
        (
            index + 1
            for row in rows
            for index, cell in enumerate(row)
            if cell.value is not None
        ),
        default=0,
    )
    rows = [_fit_row(row, width) for row in rows] or [()]
    return (
        _ExcelSheet(worksheet.title, rows),
        empty_rows_after_data,
        file_width - width,
    )


def _fit_row(row: Tuple, width: int) -> Tuple:
    """Cut or pad `row` with empty cells to exactly `width` cells."""
    return tuple(row[:width]) + (EMPTY_CELL,) * (width - len(row[:width]))


def _read_workbook(excel_file_path: str) -> Tuple[List[_ExcelSheet], List[str]]:
    """
    Read every sheet of the Excel file into memory (see _ExcelSheet).

    Returns the sheets and the warnings about them (see ExcelSurveyDefinitionReader).
    """
    sheets: List[_ExcelSheet] = []
    warnings: List[str] = []
    workbook = openpyxl.load_workbook(excel_file_path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            sheet, empty_rows_after_data, empty_columns_after_data = _read_sheet(
                worksheet
            )
            sheets.append(sheet)
            if empty_rows_after_data >= EMPTY_ROWS_WARNING_THRESHOLD:
                warnings.append(
                    f"{sheet.title} sheet has {empty_rows_after_data} empty "
                    f"rows after its last row with data (row {len(sheet.rows)}). "
                    "Delete them in Excel to make the file faster to open."
                )
            if empty_columns_after_data >= EMPTY_COLUMNS_WARNING_THRESHOLD:
                last_column = len(sheet.rows[0])
                warnings.append(
                    f"{sheet.title} sheet has {empty_columns_after_data} "
                    "empty columns after its last column with data (column "
                    f"{get_column_letter(last_column) if last_column else 'none'}). "
                    "Delete them in Excel to make the file faster to open."
                )
    finally:
        # A read-only workbook keeps the file open until closed.
        workbook.close()
    return sheets, warnings
