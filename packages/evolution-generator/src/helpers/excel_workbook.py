# Copyright 2026, Polytechnique Montreal and contributors
# This file is licensed under the MIT License.
# License text available at https://opensource.org/licenses/MIT

# Note: This module reads the Generator Excel file once, so every script works on the
# same in-memory copy instead of re-opening the file.
from typing import List, Tuple

import openpyxl
from openpyxl.cell.read_only import EMPTY_CELL
from openpyxl.utils import get_column_letter

# Number of empty rows after the last row with data, and of empty columns after the
# last column with data, from which a sheet gets a warning. Excel keeps such rows and
# columns when a cell outside the data is formatted, and they make the file slow to
# open for every tool.
EMPTY_ROWS_WARNING_THRESHOLD = 1000
EMPTY_COLUMNS_WARNING_THRESHOLD = 100


class ExcelSheet:
    """
    One sheet read into memory, cut to the area that has data: rows after the last row
    with a value, and columns after the last column with a value, are dropped. Every
    row has the same number of cells, each with a `.value` (None when empty).
    """

    def __init__(self, title: str, rows: List[Tuple]):
        self.title = title
        self.rows = rows


class ExcelWorkbook:
    """
    The sheets of an Excel file read into memory, accessed like an openpyxl workbook
    (`workbook["Widgets"]`, `workbook.sheetnames`).

    `warnings` lists the problems found while reading that don't stop the generation,
    e.g. a sheet with many empty rows or columns after its data.
    """

    def __init__(self, sheets: List[ExcelSheet], warnings: List[str]):
        self._sheets = {sheet.title: sheet for sheet in sheets}
        self.warnings = warnings

    @property
    def sheetnames(self) -> List[str]:
        return list(self._sheets)

    def __getitem__(self, sheet_name: str) -> ExcelSheet:
        if sheet_name not in self._sheets:
            # Same error as openpyxl, so a missing sheet still says which one.
            raise KeyError(f"Worksheet {sheet_name} does not exist.")
        return self._sheets[sheet_name]


def _has_value(row: Tuple) -> bool:
    return any(cell.value is not None for cell in row)


def _read_sheet(worksheet) -> Tuple[ExcelSheet, int, int]:
    """
    Read a read-only worksheet into an ExcelSheet.

    Returns the sheet and the number of empty rows and empty columns found after its
    last row and last column with data.
    """
    # The dimension saved in the file covers every formatted cell, so reading up to it
    # can mean a million empty rows: read only the rows the file actually has instead.
    worksheet.reset_dimensions()
    rows: List[Tuple] = []
    empty_rows_after_data = 0
    # Each row read ends at its last cell in the file, with or without a value.
    file_width = 0
    for row in worksheet.iter_rows():
        file_width = max(file_width, len(row))
        if _has_value(row):
            # Keep the empty rows between rows with data: they are part of the sheet.
            rows.extend([()] * empty_rows_after_data)
            empty_rows_after_data = 0
            rows.append(row)
        else:
            empty_rows_after_data += 1

    width = max(
        (
            index + 1
            for row in rows
            for index, cell in enumerate(row)
            if cell.value is not None
        ),
        default=0,
    )
    rows = [_fit_row(row, width) for row in rows] or [()]
    return (
        ExcelSheet(worksheet.title, rows),
        empty_rows_after_data,
        file_width - width,
    )


def _fit_row(row: Tuple, width: int) -> Tuple:
    """Cut or pad `row` with empty cells to exactly `width` cells."""
    return tuple(row[:width]) + (EMPTY_CELL,) * (width - len(row[:width]))


def load_excel_workbook(excel_file_path: str) -> ExcelWorkbook:
    """Read every sheet of the Excel file into memory (see ExcelSheet)."""
    sheets: List[ExcelSheet] = []
    warnings: List[str] = []
    workbook = openpyxl.load_workbook(excel_file_path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            sheet, empty_rows_after_data, empty_columns_after_data = _read_sheet(
                worksheet
            )
            sheets.append(sheet)
            if empty_rows_after_data >= EMPTY_ROWS_WARNING_THRESHOLD:
                warnings.append(
                    f"{sheet.title} sheet has {empty_rows_after_data} empty "
                    f"rows after its last row with data (row {len(sheet.rows)}). "
                    "Delete them in Excel to make the file faster to open."
                )
            if empty_columns_after_data >= EMPTY_COLUMNS_WARNING_THRESHOLD:
                last_column = len(sheet.rows[0])
                warnings.append(
                    f"{sheet.title} sheet has {empty_columns_after_data} "
                    "empty columns after its last column with data (column "
                    f"{get_column_letter(last_column) if last_column else 'none'}). "
                    "Delete them in Excel to make the file faster to open."
                )
    finally:
        # A read-only workbook keeps the file open until closed.
        workbook.close()
    return ExcelWorkbook(sheets, warnings)
